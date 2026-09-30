from __future__ import annotations

import argparse, json, re, hashlib
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

def norm(x:Any)->str:
    if x is None:
        return ""
    s=str(x).strip().lower()
    s=re.sub(r"[^a-z0-9]+","_",s).strip("_")
    return s

def sheet_rows(ws):
    vals=list(ws.iter_rows(values_only=True))
    if not vals:
        return [],[]
    # Find the most plausible header row among first 20 rows.
    best_i=0; best_score=-1
    keywords=("target","chassis","host","substrate","product","reaction","ec","pathway","literature","doi","pmid")
    for i,row in enumerate(vals[:20]):
        cells=[norm(x) for x in row if x is not None and str(x).strip()]
        score=sum(any(k in c for k in keywords) for c in cells)
        score += min(len(cells),10)*0.05
        if score>best_score:
            best_score=score; best_i=i
    headers=[str(x).strip() if x is not None else "" for x in vals[best_i]]
    rows=[]
    for raw in vals[best_i+1:]:
        if not any(x is not None and str(x).strip() for x in raw):
            continue
        rows.append({headers[j] if headers[j] else f"column_{j+1}": raw[j] if j<len(raw) else None
                     for j in range(len(headers))})
    return headers,rows

def classify(headers):
    nh=[norm(x) for x in headers]
    joined=" ".join(nh)
    score=0
    for k,w in [
        ("pathway",3),("target",3),("chassis",3),("host",2),
        ("substrate",3),("product",3),("reaction",2),("ec",3),
        ("literature",2),("doi",1),("pmid",1)
    ]:
        if k in joined: score+=w
    return score

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--xlsx",required=True)
    ap.add_argument("--output-dir",required=True)
    args=ap.parse_args()

    xlsx=Path(args.xlsx)
    out=Path(args.output_dir); out.mkdir(parents=True,exist_ok=True)
    wb=load_workbook(xlsx,read_only=True,data_only=True)

    inventory=[]
    extracted={}
    for ws in wb.worksheets:
        headers,rows=sheet_rows(ws)
        score=classify(headers)
        inventory.append({
            "sheet":ws.title,
            "max_row":ws.max_row,
            "max_column":ws.max_column,
            "header_count":len(headers),
            "headers":headers,
            "literature_schema_score":score,
            "data_row_count":len(rows),
        })
        extracted[ws.title]=rows

    ranked=sorted(inventory,key=lambda x:(x["literature_schema_score"],x["data_row_count"]),reverse=True)
    candidate_sheets=[x for x in ranked if x["literature_schema_score"]>0][:5]

    raw_json=out/"galaxy_dataset2_raw.json"
    raw_json.write_text(json.dumps({
        "schema":"synbiocrow.galaxy_dataset2_raw.v1",
        "source":{
            "doi":"10.1038/s41467-022-32661-x",
            "supplement":"Supplementary Dataset 2",
            "file_name":xlsx.name,
            "sha256":hashlib.sha256(xlsx.read_bytes()).hexdigest(),
        },
        "sheets":extracted,
    },indent=2,sort_keys=True,default=str)+"\n")

    report={
        "schema":"synbiocrow.galaxy_dataset2_inventory.v1",
        "source_file":xlsx.name,
        "sha256":hashlib.sha256(xlsx.read_bytes()).hexdigest(),
        "sheet_count":len(inventory),
        "inventory":inventory,
        "candidate_literature_sheets":candidate_sheets,
        "truth_accessed":True,
        "scoring_performed":False,
        "note":"This run inventories published benchmark truth only. It does not generate or score SynBioCrow predictions."
    }
    (out/"galaxy_dataset2_inventory.json").write_text(json.dumps(report,indent=2,sort_keys=True)+"\n")
    print(json.dumps(report,indent=2,sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
