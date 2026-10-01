"""Recover the frozen 2.3 Galaxy target-identity blacklist exactly.

Accepts either:
- the normalized 2.3 Galaxy benchmark JSON, or
- Galaxy-SynBioCAD Supplementary Dataset 2 XLSX.

The split is reconstructed from the released v2.3.0 SHA-256 pathway-ID policy.
Target structures are canonicalized with RDKit using the same InChI-to-SMILES
logic as the released 2.3 normalization script.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
import hashlib
import json
import re
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

from .historical23_blacklist import reconstruct_split


@dataclass(frozen=True)
class HistoricalTarget:
    pathway_id: str
    target_name: str
    target_inchi: str
    target_smiles: str
    split: str


@dataclass(frozen=True)
class TargetBlacklist:
    source_kind: str
    source_sha256: str
    heldout_pathway_count: int
    heldout_unique_target_count: int
    heldout_targets: Tuple[HistoricalTarget, ...]

    def canonical_json(self) -> str:
        return json.dumps(
            {
                "source_kind": self.source_kind,
                "source_sha256": self.source_sha256,
                "heldout_pathway_count": self.heldout_pathway_count,
                "heldout_unique_target_count": self.heldout_unique_target_count,
                "heldout_targets": [asdict(x) for x in self.heldout_targets],
            },
            sort_keys=True,
            separators=(",", ":"),
        )

    def sha256(self) -> str:
        return hashlib.sha256(self.canonical_json().encode("utf-8")).hexdigest()


def sha256_file(path: str) -> str:
    h=hashlib.sha256()
    with open(path,"rb") as fh:
        for block in iter(lambda: fh.read(1024*1024),b""):
            h.update(block)
    return h.hexdigest()


def canonical_smiles_from_smiles(value: str) -> str:
    from rdkit import Chem
    mol=Chem.MolFromSmiles(str(value))
    if mol is None:
        raise ValueError(f"cannot parse SMILES: {value!r}")
    return Chem.MolToSmiles(mol,canonical=True,isomericSmiles=True)


def canonical_smiles_from_inchi(value: str) -> str:
    from rdkit import Chem
    mol=Chem.MolFromInchi(str(value),sanitize=True,removeHs=True)
    if mol is None:
        raise ValueError(f"cannot parse InChI: {value!r}")
    return Chem.MolToSmiles(mol,canonical=True,isomericSmiles=True)


def _repair_known_source_inchi(value: str) -> Optional[str]:
    """Apply only narrow, auditable repairs for malformed Dataset 2 InChI text.

    Returns a repaired string or None.  No generic chemistry guessing is done.
    """
    s=str(value).strip()
    # Dataset 2 contains an impossible terminal hydrogen count in one target:
    # InChI=1S/C5H12O/c1-5(2)3-4-6/h5-6H,3-4H2,1-2H11
    # For C5H12O the terminal methyl layer should be 1-2H3.
    bad="InChI=1S/C5H12O/c1-5(2)3-4-6/h5-6H,3-4H2,1-2H11"
    if s==bad:
        return "InChI=1S/C5H12O/c1-5(2)3-4-6/h5-6H,3-4H2,1-2H3"
    return None


def _norm_header(x) -> str:
    if x is None:
        return ""
    s=str(x).strip().lower()
    return re.sub(r"[^a-z0-9]+","_",s).strip("_")


def _sheet_rows(ws):
    vals=list(ws.iter_rows(values_only=True))
    if not vals:
        return [],[]
    best_i=0
    best_score=-1
    keywords=("target","chassis","host","substrate","product","reaction","ec","pathway","literature","doi","pmid")
    for i,row in enumerate(vals[:20]):
        cells=[_norm_header(x) for x in row if x is not None and str(x).strip()]
        score=sum(any(k in c for k in keywords) for c in cells)
        score += min(len(cells),10)*0.05
        if score>best_score:
            best_score=score
            best_i=i
    headers=[str(x).strip() if x is not None else "" for x in vals[best_i]]
    rows=[]
    for raw in vals[best_i+1:]:
        if not any(x is not None and str(x).strip() for x in raw):
            continue
        rows.append({
            headers[j] if headers[j] else f"column_{j+1}":
            raw[j] if j<len(raw) else None
            for j in range(len(headers))
        })
    return headers,rows


def records_from_normalized_json(path: str) -> List[HistoricalTarget]:
    obj=json.loads(Path(path).read_text(encoding="utf-8"))
    rows=[]
    for rec in obj.get("records",[]):
        smi=rec.get("target_smiles")
        if not smi:
            continue
        rows.append(HistoricalTarget(
            pathway_id=str(rec["pathway_id"]),
            target_name=str(rec["target_name"]).strip(),
            target_inchi=str(rec.get("target_inchi") or ""),
            target_smiles=canonical_smiles_from_smiles(smi),
            split=str(rec.get("split") or ""),
        ))
    return rows


def records_from_dataset2_xlsx(path: str) -> List[HistoricalTarget]:
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True)
    sheet=None
    for ws in wb.worksheets:
        headers,rows=_sheet_rows(ws)
        normalized={_norm_header(h):h for h in headers}
        if "pathway_lit_id" in normalized and "target_structure" in normalized:
            sheet=(ws.title,rows)
            if ws.title.strip().lower()=="literature_pathway":
                break
    if sheet is None:
        raise ValueError("could not locate literature_pathway-compatible sheet")
    _,rows=sheet
    by_pid={}
    for row in rows:
        pid=str(row.get("pathway_lit_ID") or row.get("pathway_lit_id") or "").strip()
        if not pid:
            continue
        by_pid.setdefault(pid,row)
    if len(by_pid)!=77:
        raise ValueError(f"expected 77 literature pathways, found {len(by_pid)}")

    development,heldout=reconstruct_split()
    development=set(development)
    heldout=set(heldout)
    out=[]
    parse_repairs=[]
    for pid,row in by_pid.items():
        inchi=str(row.get("target_structure") or "").strip()
        if not inchi:
            raise ValueError(f"{pid}: missing target_structure")
        try:
            smi=canonical_smiles_from_inchi(inchi)
        except ValueError:
            repaired=_repair_known_source_inchi(inchi)
            if repaired is None:
                raise
            smi=canonical_smiles_from_inchi(repaired)
            parse_repairs.append({
                "pathway_id":pid,
                "target_name":str(row.get("target_name") or "").strip(),
                "original_inchi":inchi,
                "repaired_inchi":repaired,
                "repair_reason":"narrow Dataset 2 source-text correction: impossible terminal H11 -> H3 for C5H12O",
            })
            inchi=repaired
        out.append(HistoricalTarget(
            pathway_id=pid,
            target_name=str(row.get("target_name") or "").strip(),
            target_inchi=inchi,
            target_smiles=smi,
            split="development" if pid in development else "benchmark",
        ))
    records=sorted(out,key=lambda x:int(x.pathway_id.split("_")[-1]))
    records_from_dataset2_xlsx.last_repairs=parse_repairs
    return records

records_from_dataset2_xlsx.last_repairs=[]


def build_blacklist(records: Sequence[HistoricalTarget], *, source_kind: str, source_sha256: str) -> TargetBlacklist:
    _,heldout_ids=reconstruct_split()
    heldout_ids=set(heldout_ids)
    all_ids={r.pathway_id for r in records}
    missing=sorted(heldout_ids-all_ids)
    if missing:
        raise ValueError(f"source missing held-out pathway IDs: {missing[:10]}")
    selected=tuple(r for r in records if r.pathway_id in heldout_ids)
    if len(selected)!=65:
        raise ValueError(f"expected 65 held-out pathway records, found {len(selected)}")
    unique=len({r.target_smiles for r in selected})
    return TargetBlacklist(
        source_kind=source_kind,
        source_sha256=source_sha256,
        heldout_pathway_count=len(selected),
        heldout_unique_target_count=unique,
        heldout_targets=selected,
    )


def resolve_tranche(tranche_json: str, blacklist: TargetBlacklist) -> dict:
    obj=json.loads(Path(tranche_json).read_text(encoding="utf-8"))
    heldout_by_smiles={}
    heldout_names={}
    for r in blacklist.heldout_targets:
        heldout_by_smiles.setdefault(r.target_smiles,[]).append(r.pathway_id)
        heldout_names.setdefault(r.target_name.casefold(),[]).append(r.pathway_id)

    resolutions=[]
    updated=[]
    for raw in obj.get("records",[]):
        rec=dict(raw)
        candidate=canonical_smiles_from_smiles(rec["normalized_target"])
        structure_hits=heldout_by_smiles.get(candidate,[])
        name_hits=heldout_names.get(str(rec["target_name"]).casefold(),[])
        if structure_hits:
            status="blocked_historical_23"
            basis="Exact canonical target-structure match to frozen 2.3 held-out pathway(s): "+", ".join(structure_hits)
        elif name_hits:
            status="pending"
            basis="Target name matches frozen 2.3 record(s) but canonical structure did not; manual mapping review required: "+", ".join(name_hits)
        else:
            status="verified_excluded"
            basis="No canonical target-structure or normalized target-name match in the exact 65-path frozen 2.3 held-out blacklist."
        rec["historical_exclusion_status"]=status
        rec["exclusion_basis"]=basis
        updated.append(rec)
        resolutions.append({
            "record_id":rec["record_id"],
            "target_name":rec["target_name"],
            "canonical_target":candidate,
            "status":status,
            "structure_hits":structure_hits,
            "name_hits":name_hits,
        })

    return {
        "tranche_id":obj.get("tranche_id"),
        "source_blacklist_sha256":blacklist.sha256(),
        "record_count":len(updated),
        "verified_excluded_count":sum(x["status"]=="verified_excluded" for x in resolutions),
        "blocked_historical_23_count":sum(x["status"]=="blocked_historical_23" for x in resolutions),
        "pending_count":sum(x["status"]=="pending" for x in resolutions),
        "resolutions":resolutions,
        "updated_tranche":{
            **obj,
            "version":"2.4.10",
            "records":updated,
        },
    }
